from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from export_results import build_results_workbook
from grading import describe_multiple_selection, grade_answer, grade_answers, questions_requiring_review


def test_multiple_answer_grading_rules() -> None:
    assert grade_answer("AB", "A") == "inconclusive"
    assert grade_answer("AB", "C") == "wrong"
    assert grade_answer("A", "A") == "correct"
    assert grade_answer(None, "A") == "wrong"


def test_questions_requiring_review_returns_all_multiple_answer_positions() -> None:
    answers = ("A", "AB", None, "CD")

    assert questions_requiring_review(answers) == (2, 4)


def test_multiple_selection_descriptions_distinguish_issue_types() -> None:
    assert describe_multiple_selection("BC", "A") == (
        "Multiple selection, all incorrect (selected options B, C)."
    )
    assert describe_multiple_selection("ABD", "A") == (
        "Multiple selection, correct option A together with incorrect options B, D."
    )


def test_workbook_highlights_inconclusive_answers_and_status() -> None:
    answer_key = tuple("A" for _ in range(50))
    answers: tuple[str | None, ...] = ("AB", "BC", "A", None, *tuple("B" for _ in range(46)))
    outcomes, score, complete = grade_answers(answers, answer_key)
    record = {
        "student_id": "H01234567",
        "crn": "12345",
        "answers": answers,
        "outcomes": outcomes,
        "score": score,
        "percentage": score * 2.0,
        "grading_status": "Complete" if complete else "Partial",
        "source": "scripts.pdf",
        "page": 1,
    }

    workbook = load_workbook(BytesIO(build_results_workbook(answer_key, [record])))
    sheet = workbook["MCQ Results"]

    assert sheet.cell(2, 2).value == "A/B"
    assert sheet.cell(2, 2).fill.fgColor.rgb.endswith("BDD7EE")
    assert sheet.cell(2, 2).comment.text == (
        "Multiple selection, correct option A together with incorrect option B."
    )
    assert sheet.cell(2, 3).fill.fgColor.rgb.endswith("BDD7EE")
    assert sheet.cell(2, 3).comment.text == "Multiple selection, all incorrect (selected options B, C)."
    assert sheet.cell(2, 4).fill.fgColor.rgb.endswith("C6EFCE")
    assert sheet.cell(2, 5).fill.fgColor.rgb.endswith("FFEB9C")
    assert sheet.cell(2, 55).value == "Partial"
    assert sheet.cell(2, 55).fill.fgColor.rgb.endswith("BDD7EE")
